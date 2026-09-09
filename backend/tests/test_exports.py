import csv
import os
from openpyxl import load_workbook
import pytest
from app.services import export_service, attendance_service, student_service, class_service, subject_service
from app.config import settings


@pytest.mark.asyncio
async def test_export_csv_and_excel():
    cls = await class_service.create_class({
        "name": "Export Class",
        "semester": 3,
        "section": "A",
    }, teacher_id="fake_teacher")
    subject = await subject_service.create_subject({
        "name": "Data Structures",
        "code": "DS301",
        "class_id": cls["id"],
    }, teacher_id="fake_teacher")
    student = await student_service.create_student({
        "roll_number": "EXP001",
        "name": "Export Student",
        "class_id": cls["id"],
        "semester": 3,
    }, teacher_id="fake_teacher")

    session = await attendance_service.create_session(
        {"class_id": cls["id"], "subject_id": subject["id"], "date": "2026-09-08"},
        teacher_id="fake_teacher",
    )
    await attendance_service.bulk_create_records(session["id"], [
        {"student_id": student["id"], "status": "present"}
    ], teacher_id="fake_teacher")

    csv_path = await export_service.export_attendance_csv("fake_teacher", class_id=cls["id"])
    assert csv_path and os.path.exists(csv_path)

    with open(csv_path, "r", encoding="utf-8") as f:
        reader = list(csv.DictReader(f))
    assert len(reader) == 1
    assert reader[0]["Roll Number"] == "EXP001"
    assert reader[0]["Class"] == "Export Class"
    assert reader[0]["Subject"] == "Data Structures"
    assert reader[0]["Status"] == "present"

    xlsx_path = await export_service.export_attendance_excel("fake_teacher", class_id=cls["id"])
    assert xlsx_path and os.path.exists(xlsx_path)

    wb = load_workbook(xlsx_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == tuple(export_service.EXPORT_HEADERS)
    assert len(rows) == 2
    assert rows[1][3] == "EXP001"
    assert rows[1][4] == "Export Student"

    for p in (csv_path, xlsx_path):
        if p and os.path.exists(p):
            os.remove(p)


@pytest.mark.asyncio
async def test_export_is_user_scoped():
    cls = await class_service.create_class({
        "name": "Scoped Export Class",
        "semester": 3,
        "section": "A",
    }, teacher_id="owner")
    subject = await subject_service.create_subject({
        "name": "Scoped Subject",
        "code": "SCOPE001",
        "class_id": cls["id"],
    }, teacher_id="owner")

    session = await attendance_service.create_session(
        {"class_id": cls["id"], "subject_id": subject["id"], "date": "2026-09-08"},
        teacher_id="owner",
    )

    # Another teacher cannot export this data
    assert await export_service.export_attendance_csv("other_teacher", class_id=cls["id"]) is None
    assert await export_service.export_attendance_excel("other_teacher", class_id=cls["id"]) is None